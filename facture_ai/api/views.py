# ══════════════════════════════════════════════════════════
#  api/views.py
# ══════════════════════════════════════════════════════════

import io
from datetime import date, datetime
from django.http import HttpResponse

from rest_framework.decorators import api_view, permission_classes, parser_classes
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from factures.models import Invoice, LigneFacture, EcritureComptable, AuditLog
from factures.services import process_facture
from authentication.models import Company


# ══════════════════════════════════════════════════════════
#  TEST API
# ══════════════════════════════════════════════════════════
@api_view(['GET'])
def test_api(request):
    return Response({"message": "API facture opérationnelle ✓", "status": "ok"})


# ══════════════════════════════════════════════════════════
#  UPLOAD FACTURE  →  POST /api/upload/
# ══════════════════════════════════════════════════════════
@api_view(['POST'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def upload_facture(request):
    fichier = request.FILES.get('file')
    if not fichier:
        return Response({"error": "Aucun fichier reçu (champ 'file' attendu)"}, status=400)

    try:
        text, data = process_facture(fichier)
    except ValueError as e:
        return Response({"error": str(e)}, status=400)
    except Exception as e:
        import traceback
        return Response({
            "error":     "Traitement OCR échoué",
            "details":   str(e),
            "traceback": traceback.format_exc()
        }, status=500)

    try:
        date_str     = data.get("date_facture") or ""
        date_facture = datetime.strptime(date_str, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        date_facture = date.today()

    def safe_float(val, default=0.0):
        try:
            return float(val or default)
        except (ValueError, TypeError):
            return default

    devise = str(data.get("devise") or "MAD").strip().upper()
    devises_valides = ["MAD", "EUR", "USD", "GBP", "CHF", "CAD", "JPY", "AED", "SAR", "CNY"]
    if devise not in devises_valides:
        devise = "MAD"

    fichier.seek(0)

    facture = Invoice.objects.create(
        file=fichier,
        owner=request.user,
        company=Company.objects.get(id=request.user.company.id),
        fournisseur=str(data.get("fournisseur") or "Inconnu"),
        date_facture=date_facture,
        total=safe_float(data.get("total")),
        tva=safe_float(data.get("tva")),
        devise=devise,
        status=Invoice.Status.PENDING,
    )

    for ligne in data.get("lignes") or []:
        LigneFacture.objects.create(
            facture=facture,
            designation=str(ligne.get("designation") or ""),
            quantite=int(ligne.get("quantite") or 1),
            prix_unitaire=safe_float(ligne.get("prix")),
            montant=safe_float(ligne.get("montant")),
        )

    AuditLog.log(
        action=AuditLog.Action.CREATED,
        invoice=facture,
        user=request.user,
        details=f"Upload OCR · Fournisseur: {facture.fournisseur} · Total: {facture.total} {devise}",
    )

    duplicates = facture.get_duplicates()
    duplicate_warning = None
    if duplicates.exists():
        ids = ", ".join([f"#{d.id}" for d in duplicates])
        AuditLog.log(
            action=AuditLog.Action.DUPLICATE,
            invoice=facture,
            user=request.user,
            details=f"⚠️ Doublon détecté avec : {ids}",
        )
        duplicate_warning = {
            "message":    f"⚠️ Doublon détecté avec {duplicates.count()} facture(s) existante(s) !",
            "duplicates": [
                {
                    "id":           d.id,
                    "fournisseur":  d.fournisseur,
                    "total":        str(d.total),
                    "date_facture": str(d.date_facture),
                    "status":       d.status,
                }
                for d in duplicates
            ],
        }

    return Response({
        "message":           "Facture traitée avec succès",
        "facture_id":        facture.id,
        "devise":            devise,
        "data":              data,
        "duplicate_warning": duplicate_warning,
    }, status=201)


# ══════════════════════════════════════════════════════════
#  VALIDER FACTURE  →  POST /api/validate/
# ══════════════════════════════════════════════════════════
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def validate_facture(request):
    facture_id = request.data.get('facture_id')
    if not facture_id:
        return Response({"error": "facture_id requis"}, status=400)

    try:
        facture = Invoice.objects.get(id=facture_id)
    except Invoice.DoesNotExist:
        return Response({"error": "Facture introuvable"}, status=404)

    facture.status = Invoice.Status.VALIDATED
    facture.save()

    EcritureComptable.objects.filter(invoice=facture).delete()

    total_ht = float(facture.total) - float(facture.tva)
    devise   = facture.devise or "MAD"

    ecritures_data = [
        {"compte": "6111",  "libelle": f"Achat – {facture.fournisseur}",                 "debit": total_ht,           "credit": 0},
        {"compte": "34552", "libelle": f"TVA récupérable ({devise})",                    "debit": float(facture.tva), "credit": 0},
        {"compte": "4411",  "libelle": f"Fournisseur – {facture.fournisseur} [{devise}]","debit": 0,                  "credit": float(facture.total)},
    ]

    created = []
    for e in ecritures_data:
        obj = EcritureComptable.objects.create(invoice=facture, **e)
        created.append(obj)

    AuditLog.log(
        action=AuditLog.Action.VALIDATED,
        invoice=facture,
        user=request.user,
        details=f"Validée via Upload · Écritures PCM générées · {devise}",
    )

    return Response({
        "message":    "Facture validée et écritures générées",
        "facture_id": facture.id,
        "devise":     devise,
        "ecritures": [
            {"id": e.id, "compte": e.compte, "libelle": e.libelle,
             "debit": float(e.debit), "credit": float(e.credit)}
            for e in created
        ],
    })


# ══════════════════════════════════════════════════════════
#  ÉCRITURES  →  GET /api/ecritures/<id>/
# ══════════════════════════════════════════════════════════
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_ecritures(request, facture_id):
    try:
        facture = Invoice.objects.get(id=facture_id)
    except Invoice.DoesNotExist:
        return Response({"error": "Facture introuvable"}, status=404)

    ecritures = EcritureComptable.objects.filter(invoice=facture)
    return Response({
        "facture_id":  facture.id,
        "fournisseur": facture.fournisseur,
        "devise":      facture.devise or "MAD",
        "ecritures": [
            {"id": e.id, "compte": e.compte, "libelle": e.libelle,
             "debit": float(e.debit), "credit": float(e.credit)}
            for e in ecritures
        ],
    })


# ══════════════════════════════════════════════════════════
#  LISTE FACTURES  →  GET /api/factures/
# ══════════════════════════════════════════════════════════
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def liste_factures(request):
    factures = Invoice.objects.filter(company=request.user.company).order_by('-id')
    data = [
        {
            "id":           f.id,
            "fournisseur":  f.fournisseur,
            "date_facture": str(f.date_facture),
            "total":        float(f.total),
            "tva":          float(f.tva),
            "devise":       f.devise or "MAD",
            "status":       f.status,
            "validee":      f.status == Invoice.Status.VALIDATED,
            "lignes_count": f.lignes.count(),
        }
        for f in factures
    ]
    return Response({"factures": data, "total": len(data)})


# ══════════════════════════════════════════════════════════
#  HELPER : génération PDF professionnelle
# ══════════════════════════════════════════════════════════
def _generate_invoice_pdf(facture):
    buffer = io.BytesIO()
    W, H   = A4
    doc    = SimpleDocTemplate(
        buffer, pagesize=A4,
        rightMargin=20*mm, leftMargin=20*mm,
        topMargin=20*mm,   bottomMargin=20*mm,
    )

    devise = facture.devise or "MAD"

    def st(name, **kw):
        return ParagraphStyle(name, **kw)

    s_logo   = st("logo",  fontSize=18, fontName="Helvetica-Bold", textColor=colors.HexColor("#6366f1"))
    s_invnum = st("invn",  fontSize=20, fontName="Helvetica-Bold", textColor=colors.HexColor("#1a1a2e"), alignment=TA_RIGHT)
    s_label  = st("lbl",   fontSize=8,  fontName="Helvetica-Bold", textColor=colors.HexColor("#9ca3af"), spaceAfter=1*mm)
    s_value  = st("val",   fontSize=11, fontName="Helvetica-Bold", textColor=colors.HexColor("#111827"), spaceAfter=1*mm)
    s_section= st("sec",   fontSize=9,  fontName="Helvetica-Bold", textColor=colors.HexColor("#6366f1"), spaceBefore=4*mm, spaceAfter=2*mm)
    s_th     = st("th",    fontSize=8,  fontName="Helvetica-Bold", textColor=colors.white)
    s_th_r   = st("thr",   fontSize=8,  fontName="Helvetica-Bold", textColor=colors.white, alignment=TA_RIGHT)
    s_th_c   = st("thc",   fontSize=8,  fontName="Helvetica-Bold", textColor=colors.white, alignment=TA_CENTER)
    s_td     = st("td",    fontSize=9,  fontName="Helvetica",      textColor=colors.HexColor("#374151"))
    s_td_r   = st("tdr",   fontSize=9,  fontName="Helvetica",      textColor=colors.HexColor("#374151"), alignment=TA_RIGHT)
    s_td_rb  = st("tdrb",  fontSize=9,  fontName="Helvetica-Bold", textColor=colors.HexColor("#111827"), alignment=TA_RIGHT)
    s_td_c   = st("tdc",   fontSize=9,  fontName="Helvetica",      textColor=colors.HexColor("#374151"), alignment=TA_CENTER)
    s_tl     = st("tl",    fontSize=10, fontName="Helvetica",      textColor=colors.HexColor("#6b7280"), alignment=TA_RIGHT)
    s_tv     = st("tv",    fontSize=10, fontName="Helvetica-Bold", textColor=colors.HexColor("#111827"), alignment=TA_RIGHT)
    s_grand  = st("grand", fontSize=14, fontName="Helvetica-Bold", textColor=colors.HexColor("#6366f1"), alignment=TA_RIGHT)
    s_footer = st("foot",  fontSize=7,  fontName="Helvetica",      textColor=colors.HexColor("#9ca3af"), alignment=TA_CENTER)
    s_email  = st("email", fontSize=10, fontName="Helvetica",      textColor=colors.HexColor("#6366f1"))

    story = []
    cw    = W - 40*mm

    hdr = Table([[Paragraph("Smart<b>Facture</b>", s_logo),
                  Paragraph(f"FACTURE #{facture.id}", s_invnum)]],
                colWidths=[cw*0.5, cw*0.5])
    hdr.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BOTTOMPADDING",(0,0),(-1,-1),4*mm)]))
    story.append(hdr)
    story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor("#6366f1"), spaceAfter=4*mm))

    company_name = facture.company.nom if facture.company else "—"
    owner        = facture.owner
    owner_name   = f"{getattr(owner,'prenom','') or ''} {getattr(owner,'nom','') or ''}".strip() or getattr(owner,'email','—')
    date_str     = str(facture.date_facture) if facture.date_facture else "—"
    created_str  = str(facture.created_at.date()) if facture.created_at else "—"

    STATUS_MAP = {
        "VALIDATED": ("Validée",    "#d1fae5", "#065f46"),
        "PENDING":   ("En attente", "#fef3c7", "#92400e"),
        "REJECTED":  ("Rejetée",    "#fee2e2", "#991b1b"),
    }
    s_lbl, s_bg, s_fg = STATUS_MAP.get(facture.status, ("—", "#f3f4f6", "#374151"))

    status_cell = Table([[Paragraph(s_lbl, ParagraphStyle("sl", fontSize=9, fontName="Helvetica-Bold",
                          textColor=colors.HexColor(s_fg), alignment=TA_CENTER))]],
                        colWidths=[28*mm],
                        style=TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor(s_bg)),
                                          ("TOPPADDING",(0,0),(-1,-1),2*mm),
                                          ("BOTTOMPADDING",(0,0),(-1,-1),2*mm)]))

    info = Table([[
        [Paragraph("FOURNISSEUR", s_label), Paragraph(facture.fournisseur or "—", s_value),
         Spacer(1,2*mm), Paragraph("ENTREPRISE", s_label), Paragraph(company_name, s_value)],
        [Paragraph("DÉPOSÉ PAR", s_label), Paragraph(owner_name, s_value),
         Spacer(1,2*mm), Paragraph("EMAIL", s_label), Paragraph(getattr(owner,'email','—'), s_email)],
        [Paragraph("DATE FACTURE", s_label), Paragraph(date_str, s_value),
         Spacer(1,2*mm), Paragraph("DATE CRÉATION", s_label), Paragraph(created_str, s_value),
         Spacer(1,2*mm), Paragraph("STATUT", s_label), status_cell],
    ]], colWidths=[cw/3, cw/3, cw/3])
    info.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),
                               ("LEFTPADDING",(0,0),(-1,-1),0),
                               ("RIGHTPADDING",(0,0),(-1,-1),4*mm),
                               ("TOPPADDING",(0,0),(-1,-1),0),
                               ("BOTTOMPADDING",(0,0),(-1,-1),0)]))
    story.append(info)
    story.append(Spacer(1,5*mm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e5e7eb"), spaceAfter=4*mm))

    story.append(Paragraph("LIGNES DE FACTURE", s_section))
    lignes = list(facture.lignes.order_by("id"))
    if lignes:
        rows = [[Paragraph("DÉSIGNATION",s_th), Paragraph("QTÉ",s_th_c),
                 Paragraph("P.U.",s_th_r), Paragraph("MONTANT",s_th_r)]]
        for l in lignes:
            rows.append([
                Paragraph(l.designation or "—", s_td),
                Paragraph(str(l.quantite), s_td_c),
                Paragraph(f"{float(l.prix_unitaire):,.2f} {devise}", s_td_r),
                Paragraph(f"{float(l.montant):,.2f} {devise}", s_td_rb),
            ])
        lt = Table(rows, colWidths=[cw*0.5, cw*0.1, cw*0.2, cw*0.2], repeatRows=1)
        lt.setStyle(TableStyle([
            ("BACKGROUND",    (0,0),(-1,0), colors.HexColor("#6366f1")),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.HexColor("#f9fafb"),colors.white]),
            ("GRID",          (0,0),(-1,-1),0.3,colors.HexColor("#e5e7eb")),
            ("TOPPADDING",    (0,0),(-1,-1),2.5*mm),
            ("BOTTOMPADDING", (0,0),(-1,-1),2.5*mm),
            ("LEFTPADDING",   (0,0),(-1,-1),2*mm),
            ("RIGHTPADDING",  (0,0),(-1,-1),2*mm),
            ("VALIGN",        (0,0),(-1,-1),"MIDDLE"),
        ]))
        story.append(lt)
    else:
        story.append(Paragraph("Aucune ligne enregistrée.", s_td))

    story.append(Spacer(1,5*mm))

    ht    = float(facture.total) - float(facture.tva)
    tva   = float(facture.tva)
    total = float(facture.total)
    totals = Table([
        [Paragraph("Montant HT :", s_tl), Paragraph(f"{ht:,.2f} {devise}", s_tv)],
        [Paragraph("TVA :",        s_tl), Paragraph(f"{tva:,.2f} {devise}", s_tv)],
        [Paragraph("TOTAL TTC :",  s_grand), Paragraph(f"{total:,.2f} {devise}", s_grand)],
    ], colWidths=[cw*0.7, cw*0.3])
    totals.setStyle(TableStyle([
        ("LINEABOVE",     (0,2),(-1,2),1.5,colors.HexColor("#6366f1")),
        ("TOPPADDING",    (0,0),(-1,-1),2*mm),
        ("BOTTOMPADDING", (0,0),(-1,-1),2*mm),
        ("LEFTPADDING",   (0,0),(-1,-1),0),
        ("RIGHTPADDING",  (0,0),(-1,-1),0),
        ("BACKGROUND",    (0,2),(-1,2),colors.HexColor("#f5f3ff")),
    ]))
    story.append(totals)

    ecritures = list(facture.ecritures.order_by("id"))
    if ecritures:
        story.append(Spacer(1,5*mm))
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e5e7eb"), spaceAfter=4*mm))
        story.append(Paragraph("ÉCRITURES COMPTABLES PCM", s_section))
        ec_rows = [[Paragraph("COMPTE",s_th), Paragraph("LIBELLÉ",s_th),
                    Paragraph("DÉBIT",s_th_r), Paragraph("CRÉDIT",s_th_r)]]
        for e in ecritures:
            ec_rows.append([
                Paragraph(e.compte, ParagraphStyle("ec1",fontSize=9,fontName="Helvetica-Bold",textColor=colors.HexColor("#6366f1"))),
                Paragraph(e.libelle, s_td),
                Paragraph(f"{float(e.debit):,.2f}" if e.debit  else "—", ParagraphStyle("ec2",fontSize=9,fontName="Helvetica",textColor=colors.HexColor("#059669"),alignment=TA_RIGHT)),
                Paragraph(f"{float(e.credit):,.2f}" if e.credit else "—", ParagraphStyle("ec3",fontSize=9,fontName="Helvetica",textColor=colors.HexColor("#dc2626"),alignment=TA_RIGHT)),
            ])
        ec = Table(ec_rows, colWidths=[cw*0.15, cw*0.5, cw*0.175, cw*0.175])
        ec.setStyle(TableStyle([
            ("BACKGROUND",    (0,0),(-1,0),colors.HexColor("#4f46e5")),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.HexColor("#f9fafb"),colors.white]),
            ("GRID",          (0,0),(-1,-1),0.3,colors.HexColor("#e5e7eb")),
            ("TOPPADDING",    (0,0),(-1,-1),2.5*mm),
            ("BOTTOMPADDING", (0,0),(-1,-1),2.5*mm),
            ("LEFTPADDING",   (0,0),(-1,-1),2*mm),
            ("RIGHTPADDING",  (0,0),(-1,-1),2*mm),
            ("VALIGN",        (0,0),(-1,-1),"MIDDLE"),
        ]))
        story.append(ec)

    story.append(Spacer(1,8*mm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e5e7eb"), spaceAfter=3*mm))
    story.append(Paragraph(
        f"Document généré par <b>SmartFacture</b> · Facture #{facture.id} · {company_name} · {devise}",
        s_footer,
    ))

    doc.build(story)
    return buffer.getvalue()


# ══════════════════════════════════════════════════════════
#  EXPORT PDF  →  GET /api/export/pdf/<id>/
# ══════════════════════════════════════════════════════════
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_pdf(request, facture_id):
    try:
        facture = Invoice.objects.get(id=facture_id)
    except Invoice.DoesNotExist:
        return Response({"error": "Facture introuvable"}, status=404)

    AuditLog.log(
        action=AuditLog.Action.EXPORTED,
        invoice=facture,
        user=request.user,
        details=f"Export PDF · {facture.fournisseur} · {facture.total} {facture.devise or 'MAD'}",
    )

    pdf_bytes = _generate_invoice_pdf(facture)
    response  = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="facture_{facture_id}.pdf"'
    return response


# ══════════════════════════════════════════════════════════
#  EXPORT EXCEL (une facture)
# ══════════════════════════════════════════════════════════
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_excel(request, facture_id):
    try:
        facture = Invoice.objects.get(id=facture_id)
    except Invoice.DoesNotExist:
        return Response({"error": "Facture introuvable"}, status=404)

    AuditLog.log(
        action=AuditLog.Action.EXPORTED,
        invoice=facture,
        user=request.user,
        details=f"Export Excel · {facture.fournisseur} · {facture.total} {facture.devise or 'MAD'}",
    )

    devise = facture.devise or "MAD"
    wb  = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Facture"
    dark = "1a1a2e"; yellow = "e8ff47"; light = "f4f3ef"

    ws1.merge_cells('A1:D1')
    ws1['A1'] = f"FACTURE #{facture.id} – {facture.fournisseur}"
    ws1['A1'].font      = Font(bold=True, size=14, color=yellow)
    ws1['A1'].fill      = PatternFill("solid", fgColor=dark)
    ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws1.row_dimensions[1].height = 30

    for i, (k, v) in enumerate([
        ("Fournisseur", facture.fournisseur),
        ("Date",        str(facture.date_facture)),
        ("Total TTC",   f"{facture.total:.2f} {devise}"),
        ("TVA",         f"{facture.tva:.2f} {devise}"),
        ("Devise",      devise),
        ("Statut",      "Validée" if facture.status == Invoice.Status.VALIDATED else "En attente"),
    ], start=2):
        ws1[f'A{i}'] = k; ws1[f'B{i}'] = v
        ws1[f'A{i}'].font      = Font(bold=True, color="FFFFFF")
        ws1[f'A{i}'].fill      = PatternFill("solid", fgColor=dark)
        ws1[f'A{i}'].alignment = Alignment(vertical='center', indent=1)
        ws1[f'B{i}'].alignment = Alignment(vertical='center', indent=1)
        if i % 2 == 0:
            ws1[f'B{i}'].fill = PatternFill("solid", fgColor="f9f9f9")

    ws1.append([])
    headers = ["Désignation", "Quantité", f"Prix unitaire ({devise})", f"Montant ({devise})"]
    ws1.append(headers)
    hr = ws1.max_row
    for col in range(1, 5):
        c = ws1.cell(row=hr, column=col)
        c.font = Font(bold=True, color=yellow); c.fill = PatternFill("solid", fgColor=dark)
        c.alignment = Alignment(horizontal='center', vertical='center')

    for l in facture.lignes.all():
        ws1.append([l.designation, l.quantite, l.prix_unitaire, l.montant])
        row = ws1.max_row
        for col in range(1, 5):
            ws1.cell(row=row, column=col).alignment = Alignment(vertical='center', indent=1)
        if row % 2 == 0:
            for col in range(1, 5):
                ws1.cell(row=row, column=col).fill = PatternFill("solid", fgColor="f9f9f9")

    tr = ws1.max_row + 1
    ws1.cell(tr, 3, "TOTAL TTC").font = Font(bold=True)
    ws1.cell(tr, 4, facture.total).font = Font(bold=True)
    for col in range(1, 5):
        ws1.cell(tr, col).fill = PatternFill("solid", fgColor=light)

    for col, w in [('A',40),('B',12),('C',18),('D',18)]:
        ws1.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Écritures comptables")
    ws2.append(["Compte", "Libellé", f"Débit ({devise})", f"Crédit ({devise})"])
    for col in range(1, 5):
        c = ws2.cell(row=1, column=col)
        c.font = Font(bold=True, color=yellow); c.fill = PatternFill("solid", fgColor=dark)
        c.alignment = Alignment(horizontal='center')

    for e in EcritureComptable.objects.filter(invoice=facture):
        ws2.append([e.compte, e.libelle,
                    e.debit  if e.debit  > 0 else "",
                    e.credit if e.credit > 0 else ""])
        row = ws2.max_row
        if e.debit  > 0: ws2.cell(row, 3).font = Font(color="dc2626")
        if e.credit > 0: ws2.cell(row, 4).font = Font(color="0a7c3e")

    for col, w in [('A',12),('B',40),('C',16),('D',16)]:
        ws2.column_dimensions[col].width = w

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="facture_{facture_id}.xlsx"'
    wb.save(response)
    return response


# ══════════════════════════════════════════════════════════
#  EXPORT EXCEL ALL
# ══════════════════════════════════════════════════════════
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_excel_all(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Toutes les factures"
    dark = "1a1a2e"; yellow = "e8ff47"

    headers = ["ID", "Fournisseur", "Date", "Total TTC", "TVA", "Devise", "Statut"]
    ws.append(headers)
    for col in range(1, 8):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color=yellow); c.fill = PatternFill("solid", fgColor=dark)
        c.alignment = Alignment(horizontal='center')

    factures = Invoice.objects.filter(company=request.user.company).order_by('-id')
    for f in factures:
        ws.append([f.id, f.fournisseur, str(f.date_facture),
                   float(f.total), float(f.tva), f.devise or "MAD",
                   "Validée" if f.status == Invoice.Status.VALIDATED else "En attente"])

    for col, w in [('A',6),('B',30),('C',14),('D',14),('E',14),('F',10),('G',14)]:
        ws.column_dimensions[col].width = w

    AuditLog.log(
        action=AuditLog.Action.EXPORTED,
        invoice=None,
        user=request.user,
        details=f"Export Excel global · {factures.count()} factures",
    )

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="toutes_factures.xlsx"'
    wb.save(response)
    return response